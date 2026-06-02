from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.core.page_ranges import parse_page_ranges
from app.services.marking import create_marked_pdf, mark_answers, parse_answer_bank
from app.services.pdf_ops import (
    add_image_watermark,
    add_text_watermark,
    image_files_to_pdf,
    merge_pdfs,
    remove_pdf_annotations,
    rotate_pdf,
    split_pdf_by_mode,
    split_pdf_by_range_text,
)
from app.services.pdf_to_word import (
    IMAGE_BASED_EDITABLE_WARNING,
    MODE_EDITABLE,
    MODE_OCR,
    MODE_PPTX,
    MODE_SIMPLE,
    OCR_NOT_ENABLED_MESSAGE,
    SCANNED_PDF_MESSAGE,
    convert_pdf_to_docx,
    find_tesseract_executable,
    get_ocr_availability,
)
from app.services.pdf_to_pptx import convert_pdf_to_pptx
from app.services.transcript import parse_transcript, transcript_to_docx, transcript_to_pdf
from app.services.watermark_removal import remove_image_watermark_rectangles
from app.services.worksheet import clean_worksheet_pdf


class WorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_page_range_parser(self) -> None:
        ranges = parse_page_ranges("1-1, 2-3, 4-end", 10)
        self.assertEqual([(r.start, r.end) for r in ranges], [(1, 1), (2, 3), (4, 10)])
        with self.assertRaises(ValueError):
            parse_page_ranges("1-3,3-4", 10)

    def test_image_to_pdf(self) -> None:
        images = [self._image(i) for i in range(3)]
        output = image_files_to_pdf(images, self.root / "images.pdf")
        self.assertEqual(len(PdfReader(str(output)).pages), 3)

    def test_image_to_pdf_respects_selected_orientation(self) -> None:
        portrait_image = self.root / "portrait.png"
        Image.new("RGB", (300, 700), "white").save(portrait_image)
        portrait = image_files_to_pdf([portrait_image], self.root / "portrait.pdf", orientation="portrait")
        landscape = image_files_to_pdf([portrait_image], self.root / "landscape.pdf", orientation="landscape")
        portrait_box = PdfReader(str(portrait)).pages[0].mediabox
        landscape_box = PdfReader(str(landscape)).pages[0].mediabox
        self.assertGreater(float(portrait_box.height), float(portrait_box.width))
        self.assertGreater(float(landscape_box.width), float(landscape_box.height))

    def test_transcript_outputs(self) -> None:
        vtt = self.root / "sample.vtt"
        vtt.write_text("WEBVTT\n\n00:00:01.000 --> 00:00:02.000\nSpeaker: Hello\n\n", encoding="utf-8")
        cues = parse_transcript(vtt, keep_timestamps=False)
        self.assertEqual(cues[0].text, "Speaker: Hello")
        docx = transcript_to_docx(vtt, self.root / "out.docx", keep_timestamps=False)
        pdf = transcript_to_pdf(vtt, self.root / "out.pdf", keep_timestamps=True)
        self.assertTrue(docx.exists())
        self.assertEqual(len(PdfReader(str(pdf)).pages), 1)

    def test_pdf_to_word_editable_layout_mode(self) -> None:
        source = self._pdf("readable.pdf", 2)
        output = self.root / "readable.docx"
        result = convert_pdf_to_docx(source, output, MODE_EDITABLE)
        self.assertTrue(result.success, result.message)
        self.assertEqual(result.pages_processed, 2)
        self.assertGreater(result.extracted_character_count, 0)
        text = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
        self.assertIn("readable.pdf page 1", text)

    def test_pdf_to_word_accepts_output_folder(self) -> None:
        source = self._pdf("folder_output.pdf", 1)
        output_folder = self.root / "word_outputs"
        output_folder.mkdir()
        result = convert_pdf_to_docx(source, output_folder, MODE_EDITABLE)
        expected = output_folder / "folder_output.docx"
        self.assertTrue(result.success, result.message)
        self.assertEqual(Path(result.output_path), expected)
        self.assertTrue(expected.exists())

    def test_pdf_to_word_simple_text_mode(self) -> None:
        source = self._pdf("simple.pdf", 2)
        output = self.root / "simple.docx"
        result = convert_pdf_to_docx(source, output, MODE_SIMPLE)
        self.assertTrue(result.success, result.message)
        text = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
        self.assertIn("simple.pdf page 1", text)
        self.assertIn("Page 2", text)

    def test_pdf_to_word_scanned_pdf_warning(self) -> None:
        source = self._image_only_pdf()
        output = self.root / "scanned.docx"
        result = convert_pdf_to_docx(source, output, MODE_SIMPLE)
        self.assertFalse(result.success)
        self.assertEqual(result.message, SCANNED_PDF_MESSAGE)
        self.assertFalse(output.exists())
        editable = convert_pdf_to_docx(source, self.root / "scanned_editable.docx", MODE_EDITABLE)
        self.assertFalse(editable.success)
        self.assertEqual(editable.message, IMAGE_BASED_EDITABLE_WARNING)

    def test_pdf_to_word_page_ranges_and_ocr_option(self) -> None:
        source = self._pdf("ranges.pdf", 3)
        output = self.root / "range.docx"
        result = convert_pdf_to_docx(source, output, MODE_SIMPLE, "2-3")
        self.assertTrue(result.success, result.message)
        self.assertEqual(result.pages_processed, 2)
        text = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
        self.assertNotIn("ranges.pdf page 1", text)
        self.assertIn("ranges.pdf page 2", text)
        ocr = convert_pdf_to_docx(source, self.root / "ocr.docx", MODE_OCR, "1-1")
        if not ocr.success:
            self.assertIn(OCR_NOT_ENABLED_MESSAGE, ocr.message)
            self.assertFalse((self.root / "ocr.docx").exists())

    def test_pdf_to_word_ocr_scanned_pdf(self) -> None:
        source = self._image_only_pdf("OCR TEST WALLACE PDFIRST")
        output = self.root / "scanned_ocr.docx"
        result = convert_pdf_to_docx(source, output, MODE_OCR, "1-end", "English")
        self.assertTrue(result.success, result.message)
        self.assertGreater(result.extracted_character_count, 0)
        text = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)
        self.assertIn("OCR TEST", text)

    def test_bundled_ocr_detection(self) -> None:
        ocr_available, message = get_ocr_availability()
        self.assertTrue(ocr_available, message)
        executable = find_tesseract_executable()
        self.assertIsNotNone(executable)
        self.assertTrue((executable.parent / "tessdata" / "eng.traineddata").exists())
        self.assertTrue((executable.parent / "tessdata" / "chi_sim.traineddata").exists())
        self.assertTrue((executable.parent / "tessdata" / "chi_tra.traineddata").exists())

    def test_pdf_to_pptx_visual_conversion(self) -> None:
        source = self._pdf("slides.pdf", 3)
        output = self.root / "slides.pptx"
        result = convert_pdf_to_pptx(source, output, "1-2")
        self.assertTrue(result.success, result.message)
        self.assertTrue(output.exists())
        self.assertGreater(result.extracted_character_count, 0)
        presentation = Presentation(str(output))
        self.assertEqual(len(presentation.slides), 2)
        first_slide = presentation.slides[0]
        pictures = [shape for shape in first_slide.shapes if shape.shape_type == 13]
        self.assertGreaterEqual(len(pictures), 1)
        self.assertEqual(first_slide.shapes[-1].shape_type, 13)
        self.assertEqual(first_slide.shapes[-1].left, 0)
        self.assertEqual(first_slide.shapes[-1].top, 0)
        self.assertEqual(first_slide.shapes[-1].width, presentation.slide_width)
        self.assertEqual(first_slide.shapes[-1].height, presentation.slide_height)
        slide_text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("slides.pdf page 1", slide_text)

    def test_pdf_merge_split_rotate(self) -> None:
        pdfs = [self._pdf(f"p{i}.pdf", i) for i in (1, 2, 3)]
        merged = merge_pdfs(pdfs, self.root / "merged.pdf")
        self.assertEqual(len(PdfReader(str(merged)).pages), 6)
        ten = self._pdf("ten.pdf", 10)
        outputs = split_pdf_by_range_text(ten, "1-1,2-3,4-10", self.root / "split")
        self.assertEqual([len(PdfReader(str(p)).pages) for p in outputs], [1, 2, 7])
        odd_outputs = split_pdf_by_mode(ten, self.root / "odd", "odd_pages")
        even_outputs = split_pdf_by_mode(ten, self.root / "even", "even_pages")
        burst_outputs = split_pdf_by_mode(ten, self.root / "burst", "burst_single_pages")
        self.assertEqual(len(odd_outputs), 5)
        self.assertEqual(len(even_outputs), 5)
        self.assertEqual(len(burst_outputs), 10)
        rotated = rotate_pdf(ten, self.root / "rotated.pdf", 90)
        self.assertEqual(PdfReader(str(rotated)).pages[0].get("/Rotate"), 90)

    def test_watermarks(self) -> None:
        base = self._pdf("base.pdf", 1)
        text_marked = add_text_watermark(base, self.root / "text_watermark.pdf", "CONFIDENTIAL")
        image = self._image(99)
        image_marked = add_image_watermark(base, self.root / "image_watermark.pdf", image)
        self.assertEqual(len(PdfReader(str(text_marked)).pages), 1)
        self.assertEqual(len(PdfReader(str(image_marked)).pages), 1)

    def test_watermark_removal_image(self) -> None:
        img = Image.new("RGB", (240, 120), "white")
        draw = ImageDraw.Draw(img)
        draw.text((30, 45), "WATERMARK", fill=(0, 0, 0))
        source = self.root / "wm.png"
        img.save(source)
        output = remove_image_watermark_rectangles(source, self.root / "clean.png", [(20, 35, 190, 75)])
        self.assertTrue(output.exists())

    def test_worksheet_cleaning_removes_annotations(self) -> None:
        source = self._pdf_with_annotation()
        cleaned = clean_worksheet_pdf(source, self.root / "cleaned.pdf", ["student name"])
        self.assertNotIn("/Annots", PdfReader(str(cleaned)).pages[0])

    def test_marking_workflow(self) -> None:
        student_pdf = self._pdf("student.pdf", 1)
        bank = self.root / "bank.json"
        bank.write_text(json.dumps({"1": "A", "2": "B"}), encoding="utf-8")
        answers = {"1": "A", "2": "C"}
        results = mark_answers(answers, parse_answer_bank(bank))
        self.assertEqual(sum(r.score for r in results), 1)
        marked = create_marked_pdf(student_pdf, results, self.root / "marked.pdf")
        self.assertEqual(len(PdfReader(str(marked)).pages), 1)

    def test_sidebar_order_and_about_version(self) -> None:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication, QComboBox, QLabel
        from app.ui.main_window import MainWindow

        app = QApplication.instance() or QApplication([])
        window = MainWindow()
        labels = [window.nav.item(i).text() for i in range(window.nav.count())]
        self.assertEqual(
            labels,
            [
                "Split PDF",
                "Merge PDF",
                "Rotate PDF",
                "PDF to Word/PPT",
                "Transcript to Word/PDF",
                "Image to PDF",
                "Add Watermark",
                "Settings",
                "About",
            ],
        )
        about = window.stack.widget(labels.index("About")).findChildren(QLabel)
        about_text = "\n".join(label.text() for label in about)
        pdf_to_word_page = window.stack.widget(labels.index("PDF to Word/PPT"))
        mode_selector = pdf_to_word_page.findChild(QComboBox, "PdfToWordMode")
        self.assertIsNotNone(mode_selector)
        self.assertEqual(mode_selector.currentText(), MODE_EDITABLE)
        self.assertEqual([mode_selector.itemText(i) for i in range(mode_selector.count())], [MODE_EDITABLE, MODE_OCR, MODE_SIMPLE, MODE_PPTX])
        ocr_available, _ = get_ocr_availability()
        mode_selector.setCurrentText(MODE_OCR)
        if not ocr_available:
            self.assertFalse(pdf_to_word_page.run_button.isEnabled())
        self.assertIn("layout-preserving PDF-to-Word/PPT conversion", about_text)
        self.assertIn("PDF to Word/PPT note:", about_text)
        self.assertIn("Version: 1.0.1 Portable", about_text)
        settings_page = window.stack.widget(labels.index("Settings"))
        language_selector = settings_page.findChild(QComboBox, "LanguageSelector")
        self.assertIsNotNone(language_selector)
        self.assertEqual([language_selector.itemText(i) for i in range(language_selector.count())], ["English", "Simplified Chinese", "Traditional Chinese"])
        window.close()
        _ = app

    def test_no_ocr_portable_hides_ocr_pdf_mode(self) -> None:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        os.environ["WALLACES_PDFIRST_NO_OCR"] = "1"
        from PySide6.QtWidgets import QApplication, QComboBox
        from app.ui.main_window import MainWindow

        app = QApplication.instance() or QApplication([])
        try:
            window = MainWindow()
            labels = [window.nav.item(i).text() for i in range(window.nav.count())]
            pdf_to_word_page = window.stack.widget(labels.index("PDF to Word/PPT"))
            mode_selector = pdf_to_word_page.findChild(QComboBox, "PdfToWordMode")
            self.assertIsNotNone(mode_selector)
            self.assertEqual([mode_selector.itemText(i) for i in range(mode_selector.count())], [MODE_EDITABLE, MODE_SIMPLE, MODE_PPTX])
            window.close()
        finally:
            os.environ.pop("WALLACES_PDFIRST_NO_OCR", None)
        _ = app

    def _pdf(self, name: str, pages: int) -> Path:
        path = self.root / name
        c = canvas.Canvas(str(path), pagesize=A4)
        for i in range(1, pages + 1):
            c.drawString(72, 760, f"{name} page {i}")
            c.showPage()
        c.save()
        return path

    def _image(self, index: int) -> Path:
        path = self.root / f"image_{index}.png"
        img = Image.new("RGB", (300, 180), (80 + index, 120, 180))
        draw = ImageDraw.Draw(img)
        draw.text((20, 20), f"Image {index}", fill="white")
        img.save(path)
        return path

    def _image_only_pdf(self, text: str = "Image 123") -> Path:
        image = self.root / "scan_image.png"
        img = Image.new("RGB", (1800, 900), "white")
        draw = ImageDraw.Draw(img)
        try:
            font = ImageFont.truetype("arial.ttf", 82)
        except Exception:
            font = ImageFont.load_default()
        draw.text((100, 180), text, fill="black", font=font)
        draw.text((100, 340), "Scanned document sample", fill="black", font=font)
        img.save(image)
        path = self.root / "image_only.pdf"
        c = canvas.Canvas(str(path), pagesize=A4)
        c.drawImage(str(image), 72, 500, width=300, height=180)
        c.showPage()
        c.save()
        return path

    def _pdf_with_annotation(self) -> Path:
        from pypdf import PdfWriter
        from pypdf.annotations import FreeText

        base = self._pdf("annotated.pdf", 1)
        reader = PdfReader(str(base))
        writer = PdfWriter()
        writer.add_page(reader.pages[0])
        writer.add_annotation(0, FreeText(text="Answer: A", rect=(100, 650, 220, 700), font="Helvetica", font_size="12pt"))
        output = self.root / "annotated_answer.pdf"
        with output.open("wb") as fh:
            writer.write(fh)
        return output


if __name__ == "__main__":
    unittest.main()
