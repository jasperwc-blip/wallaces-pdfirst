from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import fitz
from pptx.dml.color import RGBColor
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

from app.core.page_ranges import parse_page_ranges
from app.services.pdf_to_word import configure_tesseract


@dataclass(frozen=True)
class PptxResult:
    success: bool
    message: str
    pages_processed: int
    extracted_character_count: int
    output_path: str


def convert_pdf_to_pptx(
    input_pdf_path: str | Path,
    output_pptx_path: str | Path,
    page_ranges: str = "1-end",
    ocr_language: str = "eng",
) -> PptxResult:
    source = Path(input_pdf_path)
    output = _pptx_path(output_pptx_path, source)
    if not source.exists():
        return PptxResult(False, f"PDF file not found: {source}", 0, 0, str(output))
    if source.suffix.lower() != ".pdf":
        return PptxResult(False, "Please select a PDF file.", 0, 0, str(output))
    if output.suffix.lower() != ".pptx":
        return PptxResult(False, "Please choose a PowerPoint .pptx output path.", 0, 0, str(output))

    try:
        pdf = fitz.open(str(source))
    except Exception as exc:
        return PptxResult(False, f"Could not open PDF. It may be encrypted or corrupted. Details: {exc}", 0, 0, str(output))

    try:
        ranges = parse_page_ranges(page_ranges or "1-end", pdf.page_count)
        pages = [page for page_range in ranges for page in range(page_range.start, page_range.end + 1)]
        if not pages:
            return PptxResult(False, "No pages selected.", 0, 0, str(output))

        first_page = pdf[pages[0] - 1]
        slide_cx, slide_cy = _slide_size(first_page.rect)
        prs = Presentation()
        prs.slide_width = slide_cx
        prs.slide_height = slide_cy
        blank_layout = prs.slide_layouts[6]
        extracted_chars = 0

        for page_number in pages:
            page = pdf[page_number - 1]
            slide = prs.slides.add_slide(blank_layout)
            extracted_chars += _add_page_content(slide, page, slide_cx, slide_cy, ocr_language, include_page_background=True)
    except Exception as exc:
        return PptxResult(False, f"Could not render PDF pages for PPTX. Details: {exc}", 0, 0, str(output))
    finally:
        pdf.close()

    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(output))
    except Exception as exc:
        if output.exists():
            output.unlink()
        return PptxResult(False, f"Could not create PPTX. Details: {exc}", 0, 0, str(output))
    return PptxResult(True, f"Converted PDF pages to editable PPTX: {output}", len(pages), extracted_chars, str(output))


def _pptx_path(path: str | Path, source: Path) -> Path:
    raw = str(path).strip()
    default_name = f"{source.stem}.pptx"
    if not raw:
        return Path.cwd() / "output" / default_name
    output = Path(raw)
    if output.exists() and output.is_dir():
        return output / default_name
    if raw.endswith(("\\", "/")):
        return output / default_name
    if not output.suffix:
        output = output.with_suffix(".pptx")
    return output


def _slide_size(rect: fitz.Rect) -> tuple[int, int]:
    width = max(float(rect.width), 1.0)
    height = max(float(rect.height), 1.0)
    base_cx = 9144000
    return base_cx, int(base_cx * height / width)


def _add_page_content(
    slide,
    page: fitz.Page,
    slide_cx: int,
    slide_cy: int,
    ocr_language: str,
    include_page_background: bool = True,
) -> int:
    page_width = max(float(page.rect.width), 1.0)
    page_height = max(float(page.rect.height), 1.0)
    scale_x = slide_cx / page_width
    scale_y = slide_cy / page_height
    page_dict = page.get_text("dict")
    chars = 0

    page_background = _render_page_png(page) if include_page_background else None
    if not include_page_background:
        for block in page_dict.get("blocks", []):
            if block.get("type") == 1 and block.get("image"):
                x0, y0, x1, y1 = block["bbox"]
                width = max(int((x1 - x0) * scale_x), 1)
                height = max(int((y1 - y0) * scale_y), 1)
                slide.shapes.add_picture(BytesIO(block["image"]), Emu(int(x0 * scale_x)), Emu(int(y0 * scale_y)), width=Emu(width), height=Emu(height))

    for block in page_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = [span for span in line.get("spans", []) if span.get("text", "").strip()]
            if not spans:
                continue
            x0, y0, x1, y1 = line["bbox"]
            left = Emu(int(x0 * scale_x))
            top = Emu(int(y0 * scale_y))
            width = Emu(max(int((x1 - x0) * scale_x), 1))
            height = Emu(max(int((y1 - y0) * scale_y * 1.25), 1))
            box = slide.shapes.add_textbox(left, top, width, height)
            frame = box.text_frame
            frame.clear()
            frame.margin_left = 0
            frame.margin_right = 0
            frame.margin_top = 0
            frame.margin_bottom = 0
            paragraph = frame.paragraphs[0]
            paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.0
            for span in spans:
                text = span.get("text", "")
                chars += len(text)
                run = paragraph.add_run()
                run.text = text
                run.font.size = Pt(max(float(span.get("size", 10)), 1))
                run.font.name = span.get("font") or "Arial"
                run.font.color.rgb = _span_color(span.get("color", 0))
                if include_page_background:
                    _set_run_transparency(run, 0)
                flags = int(span.get("flags", 0))
                run.font.italic = bool(flags & 2)
                run.font.bold = bool(flags & 16)
    if chars == 0:
        chars = _add_ocr_content(slide, page, slide_cx, slide_cy, ocr_language, include_page_background=not include_page_background)
    if page_background:
        slide.shapes.add_picture(BytesIO(page_background), 0, 0, width=slide_cx, height=slide_cy)
    return chars


def _add_ocr_content(
    slide,
    page: fitz.Page,
    slide_cx: int,
    slide_cy: int,
    ocr_language: str,
    include_page_background: bool = True,
) -> int:
    try:
        import pytesseract
        from PIL import Image
    except Exception:
        return 0

    tesseract = configure_tesseract(pytesseract)
    if not tesseract:
        return 0

    image_bytes = _render_page_png(page)
    if include_page_background:
        slide.shapes.add_picture(BytesIO(image_bytes), 0, 0, width=slide_cx, height=slide_cy)

    image = Image.open(BytesIO(image_bytes))
    data = pytesseract.image_to_data(image, lang=_ocr_language_code(ocr_language), output_type=pytesseract.Output.DICT)
    scale_x = slide_cx / max(float(image.width), 1.0)
    scale_y = slide_cy / max(float(image.height), 1.0)
    chars = 0
    for index, text in enumerate(data.get("text", [])):
        clean = str(text).strip()
        if not clean:
            continue
        try:
            confidence = float(data.get("conf", ["-1"])[index])
        except Exception:
            confidence = -1
        if confidence < 25:
            continue
        left = Emu(int(float(data["left"][index]) * scale_x))
        top = Emu(int(float(data["top"][index]) * scale_y))
        width = Emu(max(int(float(data["width"][index]) * scale_x * 1.25), 1))
        height = Emu(max(int(float(data["height"][index]) * scale_y * 1.35), 1))
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.clear()
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = clean
        run.font.size = Pt(max(float(data["height"][index]) / 2.2, 6))
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(0, 0, 0)
        _set_run_transparency(run, 0)
        chars += len(clean)
    return chars


def _render_page_png(page: fitz.Page) -> bytes:
    zoom = 2.0
    pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False, annots=True)
    return pixmap.tobytes("png")


def _set_run_transparency(run, alpha: int) -> None:
    """Make the editable text layer invisible while keeping text selectable."""
    rpr = run.font._element
    solid_fill = rpr.find(qn("a:solidFill"))
    if solid_fill is None:
        run.font.color.rgb = RGBColor(0, 0, 0)
        solid_fill = rpr.find(qn("a:solidFill"))
    if solid_fill is None:
        return
    for child in list(solid_fill):
        if child.tag.endswith("srgbClr") or child.tag.endswith("schemeClr"):
            child.append(parse_xml(f'<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="{alpha}"/>'))
            return


def _span_color(color: int) -> RGBColor:
    red = (int(color) >> 16) & 255
    green = (int(color) >> 8) & 255
    blue = int(color) & 255
    return RGBColor(red, green, blue)


def _ocr_language_code(label: str) -> str:
    return {
        "English": "eng",
        "Simplified Chinese": "chi_sim",
        "Traditional Chinese": "chi_tra",
        "English + Simplified Chinese": "eng+chi_sim",
        "English + Traditional Chinese": "eng+chi_tra",
        "eng": "eng",
        "chi_sim": "chi_sim",
        "chi_tra": "chi_tra",
        "eng+chi_sim": "eng+chi_sim",
        "eng+chi_tra": "eng+chi_tra",
    }.get(label, "eng")
