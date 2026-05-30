from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


TIMESTAMP_RE = re.compile(r"^\d{1,2}:?\d{2}:\d{2}[\.,]\d{3}\s+-->\s+\d{1,2}:?\d{2}:\d{2}[\.,]\d{3}")


@dataclass
class TranscriptCue:
    timestamp: str
    text: str


def parse_transcript(path: str | Path, keep_timestamps: bool = False, clean_duplicates: bool = True) -> list[TranscriptCue]:
    source = Path(path)
    content = source.read_text(encoding="utf-8-sig", errors="replace")
    suffix = source.suffix.lower()
    if suffix == ".vtt":
        cues = _parse_vtt_or_srt(content)
    elif suffix == ".srt":
        cues = _parse_vtt_or_srt(content)
    else:
        cues = [TranscriptCue("", line.strip()) for line in content.splitlines() if line.strip()]
    if clean_duplicates:
        cues = _dedupe(cues)
    if not keep_timestamps:
        cues = [TranscriptCue("", cue.text) for cue in cues]
    return cues


def transcript_to_docx(
    input_file: str | Path,
    output_docx: str | Path,
    keep_timestamps: bool = False,
    keep_speakers: bool = True,
    clean_duplicates: bool = True,
    paragraph_grouping: bool = True,
) -> Path:
    cues = parse_transcript(input_file, keep_timestamps, clean_duplicates)
    output = Path(output_docx)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Transcript", 0)
    for paragraph in _group_cues(cues, paragraph_grouping):
        if not keep_speakers:
            paragraph = _strip_speaker(paragraph)
        doc.add_paragraph(paragraph)
    doc.save(output)
    return output


def transcript_to_pdf(
    input_file: str | Path,
    output_pdf: str | Path,
    keep_timestamps: bool = False,
    keep_speakers: bool = True,
    clean_duplicates: bool = True,
    paragraph_grouping: bool = True,
) -> Path:
    cues = parse_transcript(input_file, keep_timestamps, clean_duplicates)
    output = Path(output_pdf)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    story = [Paragraph("Transcript", styles["Title"]), Spacer(1, 12)]
    for paragraph in _group_cues(cues, paragraph_grouping):
        if not keep_speakers:
            paragraph = _strip_speaker(paragraph)
        story.append(Paragraph(_escape(paragraph), styles["BodyText"]))
        story.append(Spacer(1, 8))
    doc.build(story)
    return output


def _parse_vtt_or_srt(content: str) -> list[TranscriptCue]:
    cues: list[TranscriptCue] = []
    lines = [line.strip("\ufeff") for line in content.splitlines()]
    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        if not line or line.upper() == "WEBVTT" or line.isdigit():
            idx += 1
            continue
        if TIMESTAMP_RE.match(line):
            timestamp = line
            idx += 1
            text_lines = []
            while idx < len(lines) and lines[idx].strip():
                text_lines.append(lines[idx].strip())
                idx += 1
            if text_lines:
                cues.append(TranscriptCue(timestamp=timestamp, text=" ".join(text_lines)))
        else:
            cues.append(TranscriptCue(timestamp="", text=line))
        idx += 1
    return cues


def _dedupe(cues: list[TranscriptCue]) -> list[TranscriptCue]:
    cleaned: list[TranscriptCue] = []
    previous = None
    for cue in cues:
        normalized = re.sub(r"\s+", " ", cue.text).strip().lower()
        if normalized and normalized != previous:
            cleaned.append(cue)
            previous = normalized
    return cleaned


def _group_cues(cues: list[TranscriptCue], enabled: bool) -> list[str]:
    if not enabled:
        return [_format_cue(cue) for cue in cues]
    paragraphs: list[str] = []
    buffer: list[str] = []
    for cue in cues:
        buffer.append(_format_cue(cue))
        if len(buffer) >= 4 or cue.text.endswith((".", "?", "!")):
            paragraphs.append(" ".join(buffer))
            buffer = []
    if buffer:
        paragraphs.append(" ".join(buffer))
    return paragraphs


def _format_cue(cue: TranscriptCue) -> str:
    return f"[{cue.timestamp}] {cue.text}" if cue.timestamp else cue.text


def _strip_speaker(text: str) -> str:
    return re.sub(r"\b[A-Z][A-Za-z ]{0,32}:\s+", "", text)


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
